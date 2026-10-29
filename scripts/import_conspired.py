#!/usr/bin/env python3
"""Normalize the official CONSPIRed XLSX splits as schema-v2 trait records.

The upstream training and testing files stay separate. The six published trait
labels are mapped by their documented vector order; the seventh CONSPIR trait
(``something_must_be_wrong``) is left unannotated because CONSPIRed does not label it.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterator

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("Missing dependency: python -m pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
TRAITS = (
    ("contradictory", "Contradictory"),
    ("overriding_suspicion", "Overriding suspicion"),
    ("nefarious_intent", "Nefarious intent"),
    ("persecuted_victim", "Persecuted victim"),
    ("immune_to_evidence", "Immune to evidence"),
    ("reinterpreting_randomness", "Re-interpreting randomness"),
)
UNLABELED_TRAIT = "something_must_be_wrong"
SOURCE_URLS = {
    "train": "https://raw.githubusercontent.com/UKPLab/conspired/main/data/context_training.xlsx",
    "test": "https://raw.githubusercontent.com/UKPLab/conspired/main/data/context_testing.xlsx",
}
SOURCE_LICENSE = "CC-BY-4.0"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _present(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _as_list(value: Any, field: str, row_number: int) -> list:
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    if not _present(value):
        return []
    if isinstance(value, str):
        try:
            parsed = ast.literal_eval(value)
        except (SyntaxError, ValueError) as exc:
            raise ValueError(f"Excel row {row_number}: invalid {field}: {value!r}") from exc
        if isinstance(parsed, (list, tuple)):
            return list(parsed)
    raise ValueError(f"Excel row {row_number}: {field} must be a list")


def iter_xlsx_rows(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    """Yield non-empty rows from the first worksheet with their Excel row numbers."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise ValueError(f"Workbook has no rows: {path}") from None
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        if "snippet" not in headers or "labels" not in headers:
            raise ValueError(f"{path} must contain snippet and labels columns")
        for excel_row, values in enumerate(rows, start=2):
            record = {header: value for header, value in zip(headers, values) if header}
            if any(_present(value) for value in record.values()):
                yield excel_row, record
    finally:
        workbook.close()


def _normalise_row(row: dict[str, Any], split: str, excel_row: int,
                   source_revision: str, context_column: str) -> tuple[dict, bool]:
    snippet = row.get("snippet")
    if not isinstance(snippet, str) or not snippet.strip():
        raise ValueError(f"{split} Excel row {excel_row}: missing snippet text")

    bits = _as_list(row.get("labels"), "labels", excel_row)
    if len(bits) != len(TRAITS) or any(value not in (0, 1, False, True) for value in bits):
        raise ValueError(f"{split} Excel row {excel_row}: labels must be six binary values")
    bits = [int(value) for value in bits]
    present_names = [display for bit, (_, display) in zip(bits, TRAITS) if bit]

    consolidated = _as_list(row.get("consolidated_trait"), "consolidated_trait", excel_row)
    if consolidated and set(consolidated) != set(present_names):
        raise ValueError(f"{split} Excel row {excel_row}: labels disagree with consolidated_trait")

    dominant_bits = _as_list(
        row.get("single_dominant_one_hot_dm"), "single_dominant_one_hot_dm", excel_row
    )
    if len(dominant_bits) != len(TRAITS) or any(value not in (0, 1, False, True) for value in dominant_bits):
        raise ValueError(f"{split} Excel row {excel_row}: dominant vector must contain six binary values")
    dominant_bits = [int(value) for value in dominant_bits]
    if sum(dominant_bits) != 1:
        raise ValueError(f"{split} Excel row {excel_row}: dominant vector must be one-hot")
    dominant_index = dominant_bits.index(1)
    if not bits[dominant_index]:
        raise ValueError(f"{split} Excel row {excel_row}: dominant trait is absent from multi-label vector")
    dominant_id, dominant_name = TRAITS[dominant_index]

    dominant_raw = _as_list(row.get("dominant_consol_trait"), "dominant_consol_trait", excel_row)
    dominant_conflict = bool(dominant_raw and dominant_raw != [dominant_name])
    annotations = {trait: ("present" if bit else "absent") for bit, (trait, _) in zip(bits, TRAITS)}
    annotations[UNLABELED_TRAIT] = "unannotated"

    doc_id = str(row.get("doc_id") or "").strip()
    source_id = str(row.get("id") or "").strip()
    begin = row.get("begin")
    end = row.get("end")
    identity = "\0".join(str(part) for part in (
        split, excel_row, doc_id, source_id, begin, end, snippet
    ))
    record_id = f"conspired-{split}-{hashlib.sha256(identity.encode('utf-8')).hexdigest()[:20]}"

    context_value = row.get(context_column) if context_column != "none" else None
    context = context_value.strip() if isinstance(context_value, str) and context_value.strip() else None
    metadata = {
        "source_excel_row": excel_row,
        "source_annotator_code": row.get("name"),
        "character_offsets": {"begin": begin, "end": end},
        "context500": row.get("context500"),
        "context1000": row.get("context1000"),
        "linkingpassage": row.get("linkingpassage"),
        "annotated_text": row.get("annotated_text"),
        "justification": row.get("Justification"),
        "confidence": row.get("Confidence"),
        "source_label_columns": {
            "labels": row.get("labels"),
            "consolidated_trait": row.get("consolidated_trait"),
            "dominant_consol_trait": row.get("dominant_consol_trait"),
            "single_dominant_one_hot_dm": row.get("single_dominant_one_hot_dm"),
            "OverallTrait": row.get("OverallTrait"),
            "DominantTrait": row.get("DominantTrait"),
            "Label": row.get("Label"),
        },
        # The upstream numeric label's meaning is not defined in the data README;
        # preserve it as provenance instead of inventing a mapping.
        "conspiracy_theory_or_mainstream_raw": row.get("ConspiracyTheoryorMainstream"),
        "dominant_label_conflict": dominant_conflict,
        "context_column_used": context_column,
    }
    metadata = {key: value for key, value in metadata.items() if value is not None}

    source_row_id = f"{source_id or 'missing-id'}:row-{excel_row}"
    record = {
        "schema_version": "2.0",
        "task_type": "trait_classification",
        "id": record_id,
        "text": snippet,
        "trait_annotations": annotations,
        "dominant_trait": dominant_id,
        "present_traits": [trait for trait, _ in TRAITS if annotations[trait] == "present"],
        "known_traits": [trait for trait, _ in TRAITS],
        "source": {
            "dataset": "ConspirED",
            "dataset_revision": source_revision,
            "source_row_id": source_row_id,
            "document_id": doc_id or source_row_id,
            "upstream_split": split,
            "license": SOURCE_LICENSE,
            "url": SOURCE_URLS[split],
        },
        "metadata": metadata,
    }
    if context:
        record["context"] = context
    return record, dominant_conflict


def import_conspired(train_xlsx: Path, test_xlsx: Path, output_dir: Path,
                     source_revision: str = "main", context_column: str = "none") -> dict:
    train_xlsx, test_xlsx, output_dir = map(Path, (train_xlsx, test_xlsx, output_dir))
    if context_column not in ("none", "context500", "context1000"):
        raise ValueError("context_column must be none, context500, or context1000")

    output_dir.mkdir(parents=True, exist_ok=True)
    counts, seen_ids, doc_ids, conflicts = Counter(), set(), {}, Counter()
    output_hashes = {}
    for split, source_path in (("train", train_xlsx), ("test", test_xlsx)):
        output_path = output_dir / f"{split}.jsonl"
        with output_path.open("w", encoding="utf-8", newline="\n") as destination:
            for excel_row, row in iter_xlsx_rows(source_path):
                record, conflict = _normalise_row(row, split, excel_row, source_revision, context_column)
                if record["id"] in seen_ids:
                    raise ValueError(f"Generated duplicate record id: {record['id']}")
                seen_ids.add(record["id"])
                doc_ids.setdefault(split, set()).add(record["source"]["document_id"])
                counts[split] += 1
                conflicts[split] += int(conflict)
                destination.write(json.dumps(record, ensure_ascii=False) + "\n")
        output_hashes[split] = sha256_file(output_path)

    overlap = doc_ids.get("train", set()) & doc_ids.get("test", set())
    if overlap:
        raise ValueError(f"Official train/test files share {len(overlap)} document IDs")

    manifest = {
        "dataset": "ConspirED",
        "dataset_revision": source_revision,
        "license": SOURCE_LICENSE,
        "source_urls": SOURCE_URLS,
        "source_sha256": {"train": sha256_file(train_xlsx), "test": sha256_file(test_xlsx)},
        "output_sha256": output_hashes,
        "upstream_splits_preserved": True,
        "train_records": counts["train"],
        "test_records": counts["test"],
        "train_documents": len(doc_ids.get("train", set())),
        "test_documents": len(doc_ids.get("test", set())),
        "train_test_document_overlap": 0,
        "dominant_consol_trait_conflicts_with_one_hot": dict(conflicts),
        "trait_vector_order": [display for _, display in TRAITS],
        "trait_mapping": {
            "six_conspired_labels": "labels vector; 0=absent and 1=present",
            UNLABELED_TRAIT: "unannotated; not present in the six-label upstream dataset",
            "dominant_trait": "single_dominant_one_hot_dm",
        },
        "context_column_used": context_column,
        "unmapped_source_field": "ConspiracyTheoryorMainstream preserved in metadata; source README does not define its numeric codes",
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return manifest


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--train-xlsx", type=Path, required=True)
    parser.add_argument("--test-xlsx", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "datasets" / "conspired")
    parser.add_argument("--source-revision", default="main")
    parser.add_argument("--context-column", choices=("none", "context500", "context1000"), default="none")
    args = parser.parse_args(argv)
    try:
        manifest = import_conspired(
            args.train_xlsx, args.test_xlsx, args.output_dir,
            args.source_revision, args.context_column,
        )
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    print(json.dumps(manifest, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
